import subprocess
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

from fcgo.resources import attachments
from fcgo.resources.attachments import extract_attachment


def test_extract_docx_text_and_table() -> None:
    document = Document()
    document.add_paragraph("项目说明")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "状态"
    table.cell(0, 1).text = "完成"
    output = BytesIO()
    document.save(output)

    result = _extract(output.getvalue(), "方案.docx")

    assert result.kind == "Word"
    assert "项目说明" in result.text
    assert "状态 | 完成" in result.text


def test_extract_xlsx_cells() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "数据"
    worksheet.append(["项目", "状态"])
    worksheet.append(["FCGO", "进行中"])
    output = BytesIO()
    workbook.save(output)

    result = _extract(output.getvalue(), "进度.xlsx")

    assert result.kind == "Excel"
    assert "[工作表：数据]" in result.text
    assert "FCGO | 进行中" in result.text


def test_extract_pptx_slide_text() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "季度计划"
    slide.placeholders[1].text = "完成内嵌文件读取"
    output = BytesIO()
    presentation.save(output)

    result = _extract(output.getvalue(), "计划.pptx")

    assert result.kind == "PowerPoint"
    assert "[幻灯片 1]" in result.text
    assert "完成内嵌文件读取" in result.text


def test_extract_text_source_without_executing_it() -> None:
    result = _extract(b"print('hello')\n", "tool.py", "text/x-python")

    assert result.kind == "文本文件"
    assert result.text == "print('hello')"


def test_extract_image_metadata_without_ocr() -> None:
    output = BytesIO()
    Image.new("RGB", (40, 30), color="white").save(output, format="PNG")

    result = _extract(output.getvalue(), "截图.png", "image/png")

    assert result.kind == "图片"
    assert "尺寸：40 × 30" in result.text
    assert "未执行 OCR" in result.note


def test_extract_image_with_optional_ocr(monkeypatch, tmp_path) -> None:
    output = BytesIO()
    Image.new("RGB", (40, 30), color="white").save(output, format="PNG")
    calls: list[list[str]] = []

    def fake_run(command, **kwargs):  # noqa: ANN001, ANN202
        calls.append(command)
        output_base = Path(command[2])
        output_base.with_suffix(".txt").write_text("识别出的文字\n", encoding="utf-8")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(attachments.subprocess, "run", fake_run)

    result = _extract(
        output.getvalue(),
        "截图.png",
        "image/png",
        ocr_enabled=True,
        ocr_command="fake-tesseract",
    )

    assert calls
    assert calls[0][0] == "fake-tesseract"
    assert "-l" in calls[0]
    assert "识别出的文字" in result.text
    assert "已执行本地 OCR" in result.note


def test_extract_image_ocr_missing_command(monkeypatch) -> None:
    output = BytesIO()
    Image.new("RGB", (40, 30), color="white").save(output, format="PNG")

    def fake_run(command, **kwargs):  # noqa: ANN001, ANN202, ARG001
        raise FileNotFoundError

    monkeypatch.setattr(attachments.subprocess, "run", fake_run)

    result = _extract(
        output.getvalue(),
        "截图.png",
        "image/png",
        ocr_enabled=True,
        ocr_command="missing-tesseract",
    )

    assert "尺寸：40 × 30" in result.text
    assert "未找到 OCR 命令：missing-tesseract" in result.note


def test_executable_is_never_parsed_as_text() -> None:
    result = _extract(b"MZ\x00\x00payload", "setup.exe", "application/octet-stream")

    assert result.kind == "可执行文件"
    assert result.text == ""
    assert "不会执行" in result.note


def _extract(
    content: bytes,
    filename: str,
    content_type: str = "application/octet-stream",
    *,
    ocr_enabled: bool = False,
    ocr_command: str = "tesseract",
):
    return extract_attachment(
        content,
        filename=filename,
        content_type=content_type,
        page_hint=None,
        max_chars=40_000,
        max_sheet_rows=200,
        max_sheet_columns=26,
        pdf_default_pages=2,
        pdf_max_pages=10,
        ocr_enabled=ocr_enabled,
        ocr_command=ocr_command,
        ocr_languages="chi_sim+eng",
        ocr_timeout_seconds=15,
        ocr_max_pixels=20_000_000,
    )
