import json
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader


@dataclass(frozen=True)
class AttachmentExtraction:
    kind: str
    text: str = ""
    note: str = ""


_TEXT_EXTENSIONS = {
    ".bat",
    ".c",
    ".cfg",
    ".conf",
    ".cpp",
    ".cs",
    ".css",
    ".csv",
    ".go",
    ".h",
    ".hpp",
    ".html",
    ".htm",
    ".ini",
    ".java",
    ".js",
    ".json",
    ".jsx",
    ".log",
    ".lua",
    ".md",
    ".php",
    ".properties",
    ".ps1",
    ".py",
    ".rb",
    ".rs",
    ".sh",
    ".sql",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
_IMAGE_EXTENSIONS = {
    ".avif",
    ".bmp",
    ".gif",
    ".heic",
    ".jpeg",
    ".jpg",
    ".png",
    ".tif",
    ".tiff",
    ".webp",
}
_AUDIO_EXTENSIONS = {".aac", ".flac", ".m4a", ".mp3", ".ogg", ".wav", ".wma"}
_VIDEO_EXTENSIONS = {
    ".avi",
    ".flv",
    ".m4v",
    ".mkv",
    ".mov",
    ".mp4",
    ".mpeg",
    ".mpg",
    ".webm",
    ".wmv",
}
_ARCHIVE_EXTENSIONS = {".7z", ".bz2", ".gz", ".jar", ".rar", ".tar", ".tgz", ".xz", ".zip"}
_EXECUTABLE_EXTENSIONS = {
    ".apk",
    ".app",
    ".com",
    ".dll",
    ".dmg",
    ".exe",
    ".iso",
    ".msi",
    ".so",
}


def extract_attachment(
    content: bytes,
    *,
    filename: str,
    content_type: str,
    page_hint: str | None,
    max_chars: int,
    max_sheet_rows: int,
    max_sheet_columns: int,
    pdf_default_pages: int,
    pdf_max_pages: int,
) -> AttachmentExtraction:
    extension = PurePath(filename).suffix.casefold()
    normalized_type = content_type.split(";", 1)[0].strip().casefold()

    if extension == ".pdf" or normalized_type == "application/pdf":
        text = _extract_pdf_pages(
            content,
            page_hint=page_hint,
            default_pages=pdf_default_pages,
            max_pages=pdf_max_pages,
        )
        if not text:
            return AttachmentExtraction(
                kind="PDF",
                note="指定页面没有可提取文本，可能是扫描版 PDF；当前版本暂不支持 OCR。",
            )
        return AttachmentExtraction(kind="PDF", text=_limit(text, max_chars))

    if extension == ".docx" or normalized_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        _preflight_office_zip(content)
        return AttachmentExtraction(
            kind="Word",
            text=_limit(_extract_docx(content), max_chars),
        )

    if extension == ".xlsx" or normalized_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        _preflight_office_zip(content)
        return AttachmentExtraction(
            kind="Excel",
            text=_limit(
                _extract_xlsx(
                    content,
                    max_rows=max_sheet_rows,
                    max_columns=max_sheet_columns,
                ),
                max_chars,
            ),
        )

    if extension == ".pptx" or normalized_type == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    ):
        _preflight_office_zip(content)
        return AttachmentExtraction(
            kind="PowerPoint",
            text=_limit(_extract_pptx(content), max_chars),
        )

    if extension in _IMAGE_EXTENSIONS or normalized_type.startswith("image/"):
        return _image_metadata(content)

    if extension in _TEXT_EXTENSIONS or _is_text_content_type(normalized_type):
        text = _decode_text(content)
        if extension in {".html", ".htm"} or normalized_type in {
            "text/html",
            "application/xhtml+xml",
        }:
            text = BeautifulSoup(text, "html.parser").get_text("\n")
        elif extension == ".json" or normalized_type == "application/json":
            text = _pretty_json(text)
        return AttachmentExtraction(kind="文本文件", text=_limit(text.strip(), max_chars))

    size = len(content)
    if extension in _AUDIO_EXTENSIONS or normalized_type.startswith("audio/"):
        return AttachmentExtraction(
            kind="音频",
            note=f"已识别音频文件（{normalized_type or '未知 MIME'}，{size} 字节）；暂不转写。",
        )
    if extension in _VIDEO_EXTENSIONS or normalized_type.startswith("video/"):
        return AttachmentExtraction(
            kind="视频",
            note=(
                f"已识别视频文件（{normalized_type or '未知 MIME'}，{size} 字节）；"
                "暂不转码或理解画面。"
            ),
        )
    if extension in _ARCHIVE_EXTENSIONS:
        return AttachmentExtraction(
            kind="压缩包",
            note=(
                f"已识别压缩文件（{normalized_type or '未知 MIME'}，{size} 字节）；"
                "出于安全原因不自动解压。"
            ),
        )
    if extension in _EXECUTABLE_EXTENSIONS:
        return AttachmentExtraction(
            kind="可执行文件",
            note=(
                f"已识别可执行或安装文件（{normalized_type or '未知 MIME'}，{size} 字节）；"
                "不会执行或分析二进制代码。"
            ),
        )
    if extension in {".doc", ".xls", ".ppt"}:
        return AttachmentExtraction(
            kind="旧版 Office",
            note=(
                f"暂不解析旧版二进制 Office 格式（{extension}，{size} 字节），"
                "建议转换为 DOCX/XLSX/PPTX。"
            ),
        )
    return AttachmentExtraction(
        kind="附件",
        note=f"暂不支持提取正文（{normalized_type or '未知 MIME'}，{size} 字节）。",
    )


def _extract_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    sections = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(values):
                sections.append(" | ".join(values))
    return "\n".join(sections)


def _extract_xlsx(content: bytes, *, max_rows: int, max_columns: int) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sections: list[str] = []
    try:
        for worksheet in workbook.worksheets[:5]:
            sections.append(f"[工作表：{worksheet.title}]")
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=max(max_rows, 1),
                max_col=max(max_columns, 1),
                values_only=True,
            ):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    sections.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(sections)


def _extract_pptx(content: bytes) -> str:
    presentation = Presentation(BytesIO(content))
    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        if index > 50:
            break
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if texts:
            sections.append(f"[幻灯片 {index}]\n" + "\n".join(texts))
    return "\n\n".join(sections)


def _image_metadata(content: bytes) -> AttachmentExtraction:
    with Image.open(BytesIO(content)) as image:
        width, height = image.size
        image_format = image.format or "未知格式"
        mode = image.mode
    return AttachmentExtraction(
        kind="图片",
        text=f"格式：{image_format}\n尺寸：{width} × {height}\n颜色模式：{mode}",
        note="当前版本未执行 OCR 或视觉内容理解。",
    )


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" not in text:
            return text
    raise ValueError("无法按支持的文本编码解析附件")


def _pretty_json(text: str) -> str:
    try:
        value = json.loads(text)
    except json.JSONDecodeError:
        return text
    return json.dumps(value, ensure_ascii=False, indent=2)


def _is_text_content_type(content_type: str) -> bool:
    return content_type.startswith("text/") or content_type in {
        "application/javascript",
        "application/json",
        "application/sql",
        "application/toml",
        "application/x-httpd-php",
        "application/x-sh",
        "application/xml",
        "application/yaml",
    }


def _preflight_office_zip(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000:
                raise ValueError("Office 文件包含过多内部条目")
            if sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
                raise ValueError("Office 文件解压后体积超过安全上限")
    except BadZipFile as exc:
        raise ValueError("Office 文件结构损坏") from exc


def _limit(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n[附件内容已截断]"


def _extract_pdf_pages(
    content: bytes,
    *,
    page_hint: str | None,
    default_pages: int,
    max_pages: int,
) -> str:
    reader = PdfReader(BytesIO(content), strict=False)
    page_count = len(reader.pages)
    selected_pages, requested_page = _selected_pdf_pages(
        page_hint,
        page_count=page_count,
        default_pages=default_pages,
        max_pages=max_pages,
    )
    if requested_page is not None and not selected_pages:
        return f"PDF 共 {page_count} 页，无法读取第 {requested_page} 页。"
    sections: list[str] = []
    for page_index in selected_pages:
        text = (reader.pages[page_index].extract_text() or "").strip()
        if text:
            sections.append(f"[PDF 第 {page_index + 1} 页]\n{text}")
    return "\n\n".join(sections)


def _selected_pdf_pages(
    page_hint: str | None,
    *,
    page_count: int,
    default_pages: int,
    max_pages: int,
) -> tuple[list[int], int | None]:
    safe_max = max(max_pages, 0)
    if page_hint and page_hint.startswith("pdf:page:"):
        requested = _positive_int(page_hint.removeprefix("pdf:page:"))
        if requested is None or requested > page_count or safe_max <= 0:
            return [], requested
        return [requested - 1], requested
    requested_first = None
    if page_hint and page_hint.startswith("pdf:first:"):
        requested_first = _positive_int(page_hint.removeprefix("pdf:first:"))
    page_limit = requested_first or max(default_pages, 0)
    page_limit = min(page_limit, safe_max, page_count)
    return list(range(page_limit)), None


def _positive_int(value: str) -> int | None:
    try:
        parsed = int(value)
    except ValueError:
        return None
    return parsed if parsed > 0 else None
